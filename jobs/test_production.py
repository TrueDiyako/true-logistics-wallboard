"""Tests for dashboard 5 (production plan) - layout copied from the real plan, week 41 2026."""
import datetime as dt, io, sys
import openpyxl
import production as p

def plan_xlsx():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    hdr = ["Date"] + [f"{k} {i}" for i in range(1, 6) for k in ("Flavor", "Foil", "Units")] + ["Comments"]
    L = wb.create_sheet("Production Plan Laudenberg")
    L.append(["Laudenberg production plan"]); L.append(hdr)
    L.append([dt.datetime(2026, 10, 2), "Sour Cola", "CZ/SK/RO/PL/HU/SI/HR/BG/BA", 10000])            # last week
    L.append([dt.datetime(2026, 10, 5), "Sour Watermelon", "CA", "35000 (all)",
              "Sour Watermelon", "CZ/SK/RO/PL/HU/SI/HR/BG/BA", 12000, None, None, None, None, None, None, None, None, None, "12000 bags extra"])
    L.append([dt.datetime(2026, 10, 6), "Caramel Popcorn (NEW)", "EN/DK/SE/NO/FI/IS (New)", 35000])
    L.append([dt.datetime(2026, 10, 7), "Sour Pineapple", "Bulk", "400kg"])                            # bulk: skipped
    B = wb.create_sheet("Production Plan Bossar")
    B.append(hdr)
    B.append([dt.datetime(2026, 10, 7), "Sour Watermelon", "CZ/SK/RO/PL/HU/SI/HR/BG/BA", 35000])
    B.append([dt.datetime(2026, 10, 10), "Mystery flavour", "XX/YY", 1200])                           # Saturday run, unknown SKU
    f = io.BytesIO(); wb.save(f); return f.getvalue()

RES = [  # real reservations (5 Oct 2026)
    {"sku": "6706", "order": "1022816", "customer": "Humble Group USA", "start": "2026-10-07", "cases": 576},
    {"sku": "6706", "order": "1022811", "customer": "Humble Group USA", "start": "2026-10-15", "cases": 2256},
    {"sku": "6206", "order": "1022847", "customer": "A1 d.o.o.", "start": "2026-10-09", "cases": 945},
    {"sku": "6206", "order": "1022823", "customer": "EUROBRANDS", "start": "2026-10-22", "cases": 1890},
    {"sku": "6206", "order": "1022848", "customer": "A1 d.o.o.", "start": "2026-10-23", "cases": 567},
]

def test_parse_week_only_and_sku_mapping():
    runs = p.parse_plan(plan_xlsx(), dt.date(2026, 10, 5))
    got = [(r["line"], r["day"], r["sku"], r["lang"], r["cases"]) for r in runs]
    assert got == [("Laudenberg", "2026-10-05", "6706", "CA", 2916), ("Laudenberg", "2026-10-05", "6206", "CZ/SK", 1000),
                   ("Laudenberg", "2026-10-06", "6310", "EN/DK", 2916), ("Bossar", "2026-10-07", "6206", "CZ/SK", 2916),
                   ("Bossar", "2026-10-10", "", "XX/YY", 100)]

def test_allocation_by_start_date_across_both_lines():
    out = p.build(plan_xlsx(), RES, dt.date(2026, 10, 5))
    L, B = out["lines"]["Laudenberg"], out["lines"]["Bossar"]
    ca = L[0]
    assert [(a["customer"], a["cases"], a["of"]) for a in ca["alloc"]] == [("Humble Group USA", 576, 576), ("Humble Group USA", 2256, 2256)]
    assert ca["free"] == 84
    cz_mon = L[1]                                   # Monday Laudenberg covers the earliest start first
    assert [(a["cases"], a["of"], a["start"]) for a in cz_mon["alloc"]] == [(945, 945, "2026-10-09"), (55, 1890, "2026-10-22")]
    cz_wed = B[0]                                   # Wednesday Bossar continues where Monday stopped
    assert [(a["cases"], a["of"]) for a in cz_wed["alloc"]] == [(1835, 1890), (567, 567)] and cz_wed["free"] == 514
    assert out["week"] == 41 and out["days"][-1] == "2026-10-10" and len(out["days"]) == 6   # Saturday shown: a run is planned

def test_max_rows_and_more():
    many = [{"sku": "6706", "order": str(i), "customer": f"C{i}", "start": f"2026-10-{10+i}", "cases": 10} for i in range(7)]
    out = p.build(plan_xlsx(), many, dt.date(2026, 10, 8))
    ca = out["lines"]["Laudenberg"][0]
    assert len(ca["alloc"]) == 4 and ca["more"] == 3 and ca["free"] == 2916 - 70

def test_not_excel_is_a_clear_error():
    class R:
        def __enter__(s): return s
        def __exit__(s, *a): pass
        def read(s): return b"<html>Sign in</html>"
    orig = p.urllib.request.urlopen; p.urllib.request.urlopen = lambda *a, **k: R()
    try:
        p.download_plan(); assert False
    except RuntimeError as e:
        assert "anyone with the link" in str(e)
    finally:
        p.urllib.request.urlopen = orig

def test_header_titles_overwritten_still_parsed():
    # real (14 Oct 2026): Laudenberg header "Date" became "9", Bossar "Foil 1" became "O"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    L = wb.create_sheet("Production Plan Laudenberg")
    L.append(["9", "Flavor 1", "Foil 1", "Units 1", "Flavor 2", "Foil 2", "Units 2", "Comments"])
    L.append([dt.datetime(2026, 10, 12), "Sour Watermelon", "US", 6000, "Sour Watermelon", "CZ/SK/RO/PL/HU/SI/HR/BG", 45000])
    B = wb.create_sheet("Production Plan Bossar")
    B.append(["Date", "Flavor 1", "O", "Units 1", "Flavor 2", "Foil 2", "Units 2"])
    B.append([dt.datetime(2026, 10, 12), "Sour Apple", "CZ/SK/RO/PL", 12000])
    f = io.BytesIO(); wb.save(f)
    runs = p.parse_plan(f.getvalue(), dt.date(2026, 10, 12))
    assert [(r["line"], r["sku"], r["cases"]) for r in runs] == [("Laudenberg", "6606", 500), ("Laudenberg", "6206", 3750),
                                                                 ("Bossar", "6207", 1000)]

def test_missing_header_is_a_clear_error_not_an_empty_week():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    L = wb.create_sheet("Production Plan Laudenberg"); L.append(["Date", "Taste 1", "Foil 1", "Units 1"])
    B = wb.create_sheet("Production Plan Bossar"); B.append(["Date", "Flavor 1", "Foil 1", "Units 1", "Flavor 2"])
    f = io.BytesIO(); wb.save(f)
    try:
        p.parse_plan(f.getvalue(), dt.date(2026, 10, 12)); assert False
    except RuntimeError as e:
        assert "Laudenberg" in str(e) and "header" in str(e)

def test_bbd_names():
    assert p.bbd_name(dt.date(2026, 10, 12)) == "12/04-28"          # real: Monday's production orders
    assert p.bbd_name(dt.date(2026, 8, 31)) == "29/02-28"           # no 31 Feb -> last day of month

class FakeTLProd:
    # real (13 Oct 2026): production orders for Monday 12 Oct, registered on Tuesday
    ORDERS = {"12/04-28%": [
        {"order_id": "1333701", "name": "12/04-28 (US)", "delete_dt": None, "create_date": "2026-10-13 10:46:58.000"},
        {"order_id": "1333706", "name": "12/04-28 (CZ/SK/RO/PL/HU/SI/HR/BA/BG)", "delete_dt": None, "create_date": "2026-10-13 11:02:00.000"},
        {"order_id": "1333799", "name": "12/04-28 (US) old", "delete_dt": "2026-10-13", "create_date": "2026-10-13"}]}
    LINES = {"1333701": [{"object_id": "3671907", "name": "6606", "output_product": "1", "unit_order_produced_f": "439.000000"},
                         {"object_id": "1", "name": "5.15.06.01", "output_product": "0", "unit_order_produced_f": "0"}],
             "1333706": [{"object_id": "3671900", "name": "6206", "output_product": "1", "unit_order_produced_f": "3600.000000"}]}
    # real batch (14 Oct): "12/04-28 (US)" made by production order 1333701: 439 made, 20 picked, 419 left
    BATCHES = {"1333701": [{"object_id": "4109384", "name": "12/04-28 (US)", "genobj_id": "3671907", "src_order_id": "1333701",
                            "unit_count_f": "439.000000", "unit_used_f": "20.000000", "unit_back_f": "419.000000"}],
               "1333706": [{"object_id": "4109390", "name": "12/04-28 (CZ/SK)", "genobj_id": "3671900", "src_order_id": "1333706",
                            "unit_count_f": "3600.000000", "unit_used_f": "1176.000000", "unit_back_f": "2424.000000"}]}
    def list_orders(self, f): return self.ORDERS.get(f["name"], [])
    def module(self, m, oid): return self.LINES.get(oid, [])
    def objects(self, module, f):
        assert module == "batch_genobj"; return self.BATCHES.get(f["src_order_id"].lstrip("="), [])

def test_actuals_from_production_orders():
    a = p.fetch_actuals(FakeTLProd(), [dt.date(2026, 10, 12)])
    assert a == {("2026-10-12", "6606"): {"cases": 439, "left": 419, "picked": 20, "batches": ["4109384"]},
                 ("2026-10-12", "6206"): {"cases": 3600, "left": 2424, "picked": 1176, "batches": ["4109390"]}}

def plan_42():
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    hdr = ["Date"] + [f"{k} {i}" for i in range(1, 6) for k in ("Flavor", "Foil", "Units")]
    L = wb.create_sheet("Production Plan Laudenberg"); L.append(hdr)
    L.append([dt.datetime(2026, 10, 12), "Sour Watermelon", "US", 6000, "Sour Watermelon", "CZ/SK/RO/PL", 45000])
    L.append([dt.datetime(2026, 10, 14), "Banana Caramel", "CZ/SK/RO/PL", 80000])
    B = wb.create_sheet("Production Plan Bossar"); B.append(hdr)
    B.append([dt.datetime(2026, 10, 12), "Sour Apple", "CZ/SK/RO/PL", 12000])
    f = io.BytesIO(); wb.save(f); return f.getvalue()

def test_past_days_show_actual_and_ready_pending_picking():
    res = [{"sku": "6606", "order": "1", "customer": "BEBIDAS", "start": "2026-11-03", "cases": 20},
           {"sku": "6206", "order": "2", "customer": "EUROBRANDS", "start": "2026-10-22", "cases": 420},
           {"sku": "6211", "order": "3", "customer": "EUROBRANDS", "start": "2026-10-22", "cases": 3402}]
    actuals = p.fetch_actuals(FakeTLProd(), [dt.date(2026, 10, 12)])
    out = p.build(plan_42(), res, dt.date(2026, 10, 13), actuals)
    L, B = out["lines"]["Laudenberg"], out["lines"]["Bossar"]
    us = L[0]
    assert us["status"] == "produced" and us["planned_cases"] == 500 and us["cases"] == 439
    # 20 picked by orders not in the open list -> green "other orders"; BEBIDAS' open 20 reserved -> red
    assert [(a["customer"], a["cases"], a["picked"], a["ready"]) for a in us["alloc"]] == \
        [("other orders", 20, True, False), ("BEBIDAS", 20, False, True)]
    assert us["free"] == 399                       # 419 left on the batch - 20 still reserved
    cz = L[1]
    assert cz["status"] == "produced" and cz["cases"] == 3600 and cz["left"] == 2424
    assert B[0]["status"] == "awaiting registration" and B[0]["cases"] == 1000 and not any(a["ready"] for a in B[0]["alloc"])
    bc = L[2]                                                   # Wednesday: still the plan
    assert bc["status"] == "planned" and bc["cases"] == 6666 and not bc["alloc"][0]["ready"]

def test_picked_orders_turn_green_and_left_comes_from_the_batch():
    # real: order line of 1323648 picked 20 x 6606 from batch 4109384 on 14 Oct
    line = {"name": "6606", "batches": '[{"batch_genobj_id":"4109384","unit_count_f":"20.000000","create_date":"2026-10-14 10:22:48"}]'}
    picks = p.picks_from_lines([("1022700", "BEBIDAS CINCO CERO SL", "2026-11-03", line),
                                ("1022701", "X", "2026-11-03", {"batches": "[]"}), ("1022702", "Y", "", {"batches": "not json"})])
    assert picks == {"4109384": [{"order": "1022700", "customer": "BEBIDAS CINCO CERO SL", "start": "2026-11-03", "cases": 20}]}
    actuals = p.fetch_actuals(FakeTLProd(), [dt.date(2026, 10, 12)])
    res = [{"sku": "6606", "order": "2", "customer": "Humble Group USA", "start": "2026-10-20", "cases": 100}]
    us = p.build(plan_42(), res, dt.date(2026, 10, 13), actuals, picks)["lines"]["Laudenberg"][0]
    assert [(a["customer"], a["cases"], a["picked"]) for a in us["alloc"]] == \
        [("BEBIDAS CINCO CERO SL", 20, True), ("Humble Group USA", 100, False)]
    assert us["alloc"][1]["ready"] and us["free"] == 319     # 419 left - 100 pending picking

def test_gum_and_mints_production_is_ignored():
    # real (15 Oct): gum/mints production orders with the same BBD naming, articles 4.13, 103, 106.16
    class TLGum(FakeTLProd):
        ORDERS = {"14/04-28%": [{"order_id": "77", "name": "14/04-28 (DK)", "delete_dt": None, "create_date": "2026-10-15 08:00:00"},
                                {"order_id": "78", "name": "14/04-28 (EN/DK)", "delete_dt": None, "create_date": "2026-10-15 09:09:08"}]}
        LINES = {"77": [{"object_id": "1", "name": "103", "output_product": "1", "unit_order_produced_f": "916"},
                        {"object_id": "2", "name": "4.13", "output_product": "1", "unit_order_produced_f": "208"},
                        {"object_id": "3", "name": "106.16", "output_product": "1", "unit_order_produced_f": "469"}],
                 "78": [{"object_id": "4", "name": "6310", "output_product": "1", "unit_order_produced_f": "6500"}]}
        BATCHES = {}
    a = p.fetch_actuals(TLGum(), [dt.date(2026, 10, 14)])
    assert list(a) == [("2026-10-14", "6310")]
    assert [p.is_dates_sku(x) for x in ("6606", "6206", "103", "4.13", "106.16", "6106", "60061")] == \
        [True, True, False, False, False, False, False]

def test_old_gum_and_mints_orders_with_the_same_bbd_are_ignored():
    # real (15 Oct): 29 production orders named "14/04-28 ...", most of them gum/mints from Jan-Oct
    class TLReal(FakeTLProd):
        ORDERS = {"14/04-28%": [
            {"order_id": "1173913", "name": "14/04-28 (AUS) True Gum Liquorice 24-case - AUS", "delete_dt": None, "create_date": "2026-02-03 10:38:18.000"},
            {"order_id": "1261361", "name": "14/04-28 (DE) True Mints Peach 18-case - DENLFRESIT", "delete_dt": None, "create_date": "2026-06-16 11:54:39.000"},
            {"order_id": "1326908", "name": "14/04-28 (UK) True Gum Liquorice 24-case - UK", "delete_dt": None, "create_date": "2026-10-01 15:04:47.000"},
            {"order_id": "1335250", "name": "14/04-28 (CZ/SK/RO/PL/HU/SI/HR/BA/BG)", "delete_dt": None, "create_date": "2026-10-15 09:01:42.000"}]}
        LINES = {"1173913": [{"object_id": "1", "name": "6604", "output_product": "1", "unit_order_produced_f": "208"}],   # even with a dates-like number
                 "1261361": [{"object_id": "2", "name": "106.16", "output_product": "1", "unit_order_produced_f": "469"}],
                 "1326908": [{"object_id": "3", "name": "4.13", "output_product": "1", "unit_order_produced_f": "153"}],
                 "1335250": [{"object_id": "4", "name": "6211", "output_product": "1", "unit_order_produced_f": "6666"}]}
        BATCHES = {}
    a = p.fetch_actuals(TLReal(), [dt.date(2026, 10, 14)])
    assert list(a) == [("2026-10-14", "6211")] and a[("2026-10-14", "6211")]["cases"] == 6666

def test_unplanned_production_is_shown():
    out = p.build(plan_42(), [], dt.date(2026, 10, 13), {("2026-10-12", "6309"): 250})
    extra = [r for r in out["lines"]["Laudenberg"] if r.get("unplanned")]
    assert len(extra) == 1 and extra[0]["cases"] == 250 and extra[0]["sku"] == "6309"

if __name__ == "__main__":
    import traceback
    fails = 0
    for n, f in list(globals().items()):
        if n.startswith("test_"):
            try: f(); print("PASS", n)
            except Exception as e: fails += 1; print("FAIL", n, repr(e)); traceback.print_exc()
    sys.exit(fails)
